"""Transcript loading, speaker classification, rendering and export helpers."""

from __future__ import annotations

import html
import io
import json
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document
from docx.shared import Pt, RGBColor

from core.config import CHAMBER, DATA_DIR

# (keywords matched against the lower-cased speaker label, icon, role, accent colour)
ROLES = [
    (("prime minister", "perdana menteri", "pm"), "🏛️", "Prime Minister", "#b8860b"),
    (("sri aman",), "🎤", "MP Sri Aman", "#1f77b4"),
    (("speaker", "yang dipertua", "yang di-pertua", "tuan yang di-pertua"), "🔔", "Speaker of the House", "#6c757d"),
]
UNIDENTIFIED = "Unidentified"
UNKNOWN_ROLE = ("👤", UNIDENTIFIED, "#4f8bf9")

SEGMENT_KEYS = ("turns", "segments", "transcript", "utterances", "results")
AUDIO_MIME = {".wav": "audio/wav", ".mp3": "audio/mpeg", ".m4a": "audio/mp4", ".mp4": "audio/mp4"}


# ---------------------------------------------------------------- data loading
def _first(d: dict, *keys: str, default: Any = None) -> Any:
    for k in keys:
        if k in d and d[k] is not None:
            return d[k]
    return default


def _parse_clock(value: str) -> float:
    """'MM:SS.ss' or 'HH:MM:SS.ss' -> seconds."""
    secs = 0.0
    for part in value.strip().split(":"):
        secs = secs * 60 + float(part)
    return secs


def _times(seg: dict) -> tuple[float, float]:
    start = _first(seg, "start", "start_time", "begin")
    end = _first(seg, "end", "end_time", "stop")
    ts = seg.get("timestamp")
    if (start is None or end is None) and isinstance(ts, str) and "-" in ts:
        try:
            ts_start, ts_end = (_parse_clock(p) for p in ts.split("-", 1))
            start = ts_start if start is None else start
            end = ts_end if end is None else end
        except ValueError:
            pass
    return float(start or 0.0), float(end or 0.0)


def _find_segment_list(raw: dict) -> list:
    for key in SEGMENT_KEYS:
        if isinstance(raw.get(key), list):
            return raw[key]
    for value in raw.values():
        if isinstance(value, list) and value and isinstance(value[0], dict) and "text" in value[0]:
            return value
    return []


def list_transcripts() -> list[Path]:
    """Transcript JSONs in data/, newest first (so fresh pipeline output is the default)."""
    files = [p for p in DATA_DIR.glob("*.json") if p.is_file()]
    return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)


@st.cache_data(show_spinner=False)
def load_transcript(path: Path, mtime: float) -> tuple[list[dict], dict]:
    """Load segments from JSON: a bare list, or a dict holding them under 'turns'/'segments'/…

    `mtime` is part of the cache key so a re-generated file is re-read.
    """
    raw = json.loads(path.read_text(encoding="utf-8"))
    meta: dict = {}
    if isinstance(raw, dict):
        meta = {k: v for k, v in raw.items() if not isinstance(v, (list, dict)) and v is not None}
        raw = _find_segment_list(raw)

    segments = []
    for i, seg in enumerate(raw):
        if not isinstance(seg, dict):
            continue
        start, end = _times(seg)
        segments.append(
            {
                "id": i,
                "start": start,
                "end": end,
                "speaker": str(_first(seg, "speaker_name", "speaker", "label", default="Unknown")),
                "text": str(_first(seg, "text", "transcript", "content", default="")).strip(),
                "confidence": _first(seg, "confidence", "similarity", "score"),
            }
        )
    return segments, meta


def resolve_audio(transcript: Path, meta: dict) -> Path | None:
    """The audio named in the JSON's `audio_file`, else `<stem without _transcript>.wav` beside it."""
    candidates = []
    if name := (meta.get("audio_file") or meta.get("Audio_File")):
        candidates.append(transcript.parent / Path(str(name)).name)
    candidates.append(transcript.parent / f"{transcript.stem.removesuffix('_transcript')}.wav")
    return next((p for p in candidates if p.is_file()), None)


@st.cache_resource(max_entries=2, show_spinner="Loading session audio…")
def load_audio_bytes(path: Path, mtime: float) -> bytes:
    """Read the audio once per file version.

    Passing a path to st.audio re-reads the whole file on every rerun (hundreds of MB for a
    full sitting); cached bytes keep seeking and filtering snappy.
    """
    return path.read_bytes()


def audio_mime(path: Path) -> str:
    return AUDIO_MIME.get(path.suffix.lower(), "audio/wav")


# ------------------------------------------------------------- presentation
def classify(speaker: str) -> tuple[str, str, str]:
    """Return (icon, role, colour) for a speaker label."""
    s = speaker.lower()
    for keywords, icon, role, colour in ROLES:
        if any(k == s or k in s.split() or (" " in k and k in s) for k in keywords):
            return icon, role, colour
    return UNKNOWN_ROLE


def fmt_ts(seconds: float) -> str:
    seconds = int(seconds)
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def highlight(text: str, query: str) -> str:
    safe = html.escape(text)
    if not query:
        return safe
    q = html.escape(query)
    lower = safe.lower()
    out, i, ql = [], 0, len(q)
    while (j := lower.find(q.lower(), i)) != -1:
        out.append(safe[i:j])
        out.append(f"<mark>{safe[j:j + ql]}</mark>")
        i = j + ql
    out.append(safe[i:])
    return "".join(out)


def pill(text: str, kind: str = "info") -> str:
    return f'<span class="pill pill-{kind}">{html.escape(text)}</span>'


def meta_row(key: str, value: str) -> str:
    return (f'<div class="meta-row"><span>{html.escape(key)}</span>'
            f'<span>{html.escape(value)}</span></div>')


def render_card(seg: dict, query: str, active: bool = False) -> str:
    icon, role, colour = classify(seg["speaker"])
    conf = ""
    if isinstance(seg["confidence"], (int, float)):
        conf = f' · sim {seg["confidence"]:.2f}'
    return f"""
<div class="card{' card-active' if active else ''}" style="--accent:{colour}">
  <div class="card-head">
    <span class="card-icon">{icon}</span>
    <span class="card-speaker">{html.escape(seg["speaker"])}</span>
    <span class="card-role">{role}{conf}</span>
    <span class="card-ts">{fmt_ts(seg["start"])} → {fmt_ts(seg["end"])}</span>
  </div>
  <div class="card-text">{highlight(seg["text"], query)}</div>
</div>"""


CSS = """
<style>
.pill {display:inline-block;padding:2px 10px;margin:2px 4px 2px 0;border-radius:999px;
       font-size:0.78rem;font-weight:600;border:1px solid transparent;white-space:nowrap}
.pill-ok   {background:rgba(40,167,69,.15);color:#28a745;border-color:rgba(40,167,69,.4)}
.pill-warn {background:rgba(255,193,7,.15);color:#c69500;border-color:rgba(255,193,7,.5)}
.pill-err  {background:rgba(220,53,69,.12);color:#dc3545;border-color:rgba(220,53,69,.4)}
.pill-info {background:rgba(31,119,180,.12);color:#1f77b4;border-color:rgba(31,119,180,.4)}
.meta-row {display:flex;justify-content:space-between;gap:8px;font-size:0.85rem;padding:2px 0}
.meta-row span:first-child {opacity:.65}
.meta-row span:last-child {text-align:right;overflow-wrap:anywhere}
.card {border-left:4px solid var(--accent);border-radius:10px;padding:10px 14px;margin:0 0 4px 0;
       background:rgba(127,127,127,.07)}
.card-active {background:rgba(255,193,7,.14);box-shadow:0 0 0 1px rgba(255,193,7,.5)}
.card-head {display:flex;align-items:center;gap:8px;margin-bottom:4px;flex-wrap:wrap}
.card-icon {font-size:1.35rem;line-height:1}
.card-speaker {font-weight:700;color:var(--accent)}
.card-role {font-size:.72rem;opacity:.7}
.card-ts {margin-left:auto;font-family:ui-monospace,monospace;font-size:.78rem;opacity:.7}
.card-text {line-height:1.55}
.card-text mark {background:#ffe066;color:#000;padding:0 2px;border-radius:3px}
</style>
"""


# ------------------------------------------------------------------ exports
def build_txt(segments: list[dict]) -> str:
    return "\n\n".join(
        f"[{fmt_ts(s['start'])} → {fmt_ts(s['end'])}] {s['speaker']}:\n{s['text']}" for s in segments
    )


@st.cache_data(show_spinner=False)
def build_docx(path: Path, mtime: float) -> bytes:
    """Render the transcript as a Hansard-style Word document (cached per file version)."""
    segments, meta = load_transcript(path, mtime)
    doc = Document()
    doc.add_heading(str(meta.get("title") or path.stem), level=0)
    for k, v in {**CHAMBER, **{k.replace("_", " ").title(): str(v) for k, v in meta.items()}}.items():
        p = doc.add_paragraph()
        p.add_run(f"{k}: ").bold = True
        p.add_run(v)

    for seg in segments:
        _, _, colour = classify(seg["speaker"])
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(10)
        head.paragraph_format.space_after = Pt(0)
        name = head.add_run(seg["speaker"])
        name.bold = True
        name.font.color.rgb = RGBColor.from_string(colour.lstrip("#"))
        ts = head.add_run(f"  [{fmt_ts(seg['start'])} → {fmt_ts(seg['end'])}]")
        ts.font.size = Pt(9)
        ts.font.color.rgb = RGBColor(0x6C, 0x75, 0x7D)
        doc.add_paragraph(seg["text"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
